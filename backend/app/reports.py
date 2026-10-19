"""Live reports built from the CURRENT snapshot and the chosen filters.
Exports never contain full Aadhaar/mobile numbers (masked columns only)."""
import csv
import io
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import calc
from .queries import att_status, employee_row, filter_employees, pay_for, sort_employees, summary

REPORTS = {
    "all_employees": "All employees",
    "group_summary": "Group-wise employee summary",
    "location_summary": "Location-wise summary",
    "daily_attendance": "Daily attendance",
    "period_attendance": "Attendance for a period",
    "employee_pay": "Employee-wise earned pay",
    "group_payroll": "Group-wise payroll and balance",
    "departed": "Departed employees",
    "recent_departures": "Recent departures",
    "missing_rates": "Missing rates",
    "data_quality": "Data-quality issues",
    "completeness": "Employee information completeness",
    "audit": "Change / audit history",
}


def _inr(v):
    return None if v is None else float(v)


def build(snap, name, f, last_sync, today):
    d = f["date"]
    lo, hi = f.get("period_from"), f.get("period_to")
    period = (lo, hi) if (lo or hi) else None
    cols, rows = [], []
    emps_f = dict(f)
    warnings = []

    def pop(status="current"):
        return sort_employees(filter_employees(snap, {**emps_f, "status": emps_f.get("status") or status}), "row", False)

    if name == "all_employees":
        cols = [("group", "Group"), ("name", "Name"), ("designation", "Designation"), ("location", "Location"), ("status", "Employment status"),
                ("att", f"Attendance {d.isoformat()}"), ("rate", "Per day pay"), ("aadhaar", "Aadhaar (masked)"), ("mobile", "Mobile (masked)")]
        rows = [employee_row(snap, e, d, period, with_pay=False) for e in pop()]
    elif name in ("group_summary", "group_payroll"):
        for g in snap.groups:
            if f.get("group") and g != f["group"]:
                continue
            emps = [e for e in filter_employees(snap, {**emps_f, "group": g, "status": emps_f.get("status") or "all"})]
            sm = summary(snap, emps, d, period, today)
            sp = snap.sheet(g)
            from .queries import group_amounts
            ga = group_amounts(snap, sp, sm)
            r = {"group": g, "total": sm["total"], "active": sm["active"], "departed": sm["departed"], "not_verified": sm["not_verified"],
                 "present": sm["attendance"]["present"], "absent": sm["attendance"]["absent"], "not_recorded": sm["attendance"]["not_recorded"],
                 "legacy_days": sm["legacy_present_days"], "confirmed": sm["pay"]["confirmed_total"], "incomplete": sm["pay"]["incomplete_count"],
                 "rate_missing": sm["pay"]["rate_missing_count"], "advance": ga["advance"]["value"], "sweets": ga["sweets"]["value"],
                 "balance": ga["balance"]["app"]}
            rows.append(r)
        if name == "group_summary":
            cols = [("group", "Group"), ("total", "Employees"), ("active", "Active"), ("departed", "Departed"), ("not_verified", "Status not verified"),
                    ("present", f"Present {d.isoformat()}"), ("absent", "Absent"), ("not_recorded", "Not recorded"), ("rate_missing", "Rate cells missing")]
        else:
            cols = [("group", "Group"), ("total", "Staff count"), ("legacy_days", "Present days (legacy)"),
                    ("rate_missing", "Rate cells missing"), ("advance", "Advance ₹ (group cell)"), ("sweets", "Sweets ₹ (group cell)"), ("balance", "Workbook formula balance (unverified)")]
            warnings.append("Compensation policy not approved: No employee has an approved payment rate or payable obligation. Figures show operational attendance days and raw source spreadsheet cells. Advance, Sweets and Balance are shared group figures and are NOT individual employee liabilities.")
    elif name == "location_summary":
        from collections import defaultdict
        agg = defaultdict(list)
        for e in pop("all"):
            agg[e["location"] or "(blank)"].append(e)
        cols = [("location", "Location"), ("total", "Employees"), ("active", "Active"), ("not_verified", "Status not verified"), ("departed", "Departed"), ("policy_status", "Compensation status")]
        for loc, es in sorted(agg.items()):
            sm = summary(snap, es, d, period, today)
            rows.append({"location": loc, "total": sm["total"], "active": sm["active"], "not_verified": sm["not_verified"], "departed": sm["departed"], "policy_status": "Rates not authorised"})
        warnings.append("Financial calculations unavailable: No compensation rates are authorised under current policy.")
    elif name == "daily_attendance":
        cols = [("group", "Group"), ("name", "Name"), ("location", "Location"), ("att", f"Attendance {d.isoformat()}")]
        rows = [employee_row(snap, e, d, None, with_pay=False) for e in pop()]
        if not snap.attendance_by_date.get(d):
            warnings.append("No attendance has been recorded for this date yet — everyone shows Not Recorded (this does NOT mean absent).")
    elif name == "period_attendance":
        lo2, hi2 = lo or d.replace(day=1), hi or d
        cols = [("group", "Group"), ("name", "Name"), ("present", "Present"), ("absent", "Absent"), ("leave", "Leave"), ("off", "Weekly off/holiday"), ("recorded", "Days recorded")]
        for e in pop():
            recs = [r for dd, r in snap.attendance.get(e["id"], {}).items() if lo2 <= dd <= hi2 and r["status"]] if e["id"] else []
            c = {k: sum(1 for r in recs if r["status"] == k) for k in ("Present", "Absent", "Leave", "Weekly Off/Holiday")}
            rows.append({"group": e["sheet"], "name": e["name_display"], "present": c["Present"], "absent": c["Absent"], "leave": c["Leave"], "off": c["Weekly Off/Holiday"], "recorded": len(recs)})
        warnings.append(f"Date-wise records only, {lo2.isoformat()} to {hi2.isoformat()}. Legacy Sept/Oct totals are in the pay reports.")
    elif name == "employee_pay":
        cols = [("group", "Group"), ("name", "Name"), ("status", "Status"), ("legacy_days", "Legacy present days"), ("rate", "Per day pay (workbook cell)"),
                ("datewise_days", "Date-wise days"), ("pay_label", "Compensation policy status")]
        rows = [employee_row(snap, e, d, period) for e in pop("all")]
        warnings.append("Compensation policy not approved: No employee has an approved compensation rate. No payment amount is calculated or authorised under current policy.")
    elif name in ("departed", "recent_departures"):
        days = f.get("recent_days") or snap.config.get("recent_departure_days") or 30
        cols = [("group", "Former group"), ("name", "Name"), ("location", "Location"), ("dep_date", "Departure date"), ("reason", "Reason"),
                ("last_att", "Last recorded attendance"), ("legacy_days", "Legacy present days"), ("pay_label", "Compensation policy status")]
        base = filter_employees(snap, {**emps_f, "status": "Departed"})
        for e in base:
            if name == "recent_departures" and not (e["dep_date"] and today - timedelta(days=days) <= e["dep_date"] <= today):
                continue
            r = employee_row(snap, e, d, period)
            recs = [dd for dd, x in snap.attendance.get(e["id"], {}).items() if x["status"]] if e["id"] else []
            r["last_att"] = max(recs).isoformat() if recs else None
            r["reason"] = e["dep_reason"]
            rows.append(r)
        warnings.append("Historical attendance and work records are preserved for departed staff. Outstanding payments are not tracked or inferred because no approved compensation ledger exists.")
        warnings.append("Outstanding/settlement amounts are not available: the workbook has no per-employee payment ledger.")
    elif name == "missing_rates":
        cols = [("group", "Group"), ("name", "Name"), ("location", "Location"), ("rate_state", "Rate state"), ("legacy_days", "Present days awaiting a rate")]
        for e in pop("all"):
            if e["rate"]["state"] in ("blank", "invalid"):
                r = employee_row(snap, e, d, period)
                r["rate_state"] = "Rate required" if e["rate"]["state"] == "blank" else "Invalid"
                rows.append(r)
    elif name == "data_quality":
        cols = [("severity", "Severity"), ("sheet", "Sheet"), ("row", "Row"), ("emp_name", "Employee"), ("message", "Issue"), ("resolved", "Acknowledged")]
        for i in snap.issues:
            if f.get("group") and i["sheet"] != f["group"]:
                continue
            rows.append({**i, "resolved": (i["resolution"] or {}).get("by")})
    elif name == "completeness":
        cols = [("group", "Group"), ("name", "Name"), ("designation", "Designation"), ("location", "Location"), ("rate", "Per day pay"),
                ("att_legacy", "Sept/Oct attendance"), ("aadhaar", "Aadhaar"), ("mobile", "Mobile"), ("from_date", "From date"), ("missing", "Missing")]
        for e in pop("all"):
            miss = [lab for lab, v in (("Designation", e["designation"]), ("Location", e["location"]), ("Rate", e["rate"]["value"]),
                                       ("Attendance", e["sept"]["value"] if e["sept"]["state"] == "number" else e["octo"]["value"]),
                                       ("Aadhaar", e["aadhaar"]), ("Mobile", e["mobile"]), ("From date", e["from_date"]["value"])) if v is None]
            yes = lambda x: "Yes" if x else "—"
            rows.append({"group": e["sheet"], "name": e["name_display"], "designation": yes(e["designation"]), "location": yes(e["location"]),
                         "rate": yes(e["rate"]["value"] is not None), "att_legacy": yes(e["sept"]["state"] == "number" or e["octo"]["state"] == "number"),
                         "aadhaar": yes(e["aadhaar"]), "mobile": yes(e["mobile"]), "from_date": yes(e["from_date"]["value"]), "missing": ", ".join(miss) or "Complete"})
    elif name == "audit":
        cols = [("ts", "When"), ("actor", "Who"), ("action", "Action"), ("emp_id", "Employee ID"), ("group", "Group"), ("field", "Field"), ("before", "Before"), ("after", "After"), ("note", "Note")]
        rows = list(reversed(snap.audit))
    else:
        return None
    unresolved = sum(1 for i in snap.issues if i["severity"] != "info" and not i["resolution"])
    meta = {"report": REPORTS[name], "generated_at": datetime.now().isoformat(timespec="seconds"), "last_sync": last_sync,
            "filters": {k: (v.isoformat() if isinstance(v, date) else v) for k, v in f.items() if v not in (None, "") and k != "date"} | {"date": d.isoformat()},
            "unresolved_warnings": unresolved, "notes": warnings, "row_count": len(rows),
            "sensitive": "Aadhaar and mobile numbers are masked in exports."}
    return {"name": name, "title": REPORTS[name], "columns": [{"key": k, "label": l} for k, l in cols], "rows": rows, "meta": meta}


def _safe(v):
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def to_csv(rep):
    buf = io.StringIO()
    w = csv.writer(buf)
    m = rep["meta"]
    for line in (f"# {m['report']}", f"# Generated: {m['generated_at']}", f"# Last synced: {m['last_sync']}",
                 f"# Filters: {m['filters']}", f"# Unresolved warnings: {m['unresolved_warnings']}", f"# {m['sensitive']}", *[f"# Note: {n}" for n in m["notes"]]):
        w.writerow([line])
    w.writerow([c["label"] for c in rep["columns"]])
    for r in rep["rows"]:
        w.writerow([_safe(r.get(c["key"])) for c in rep["columns"]])
    return "\ufeff" + buf.getvalue()


def to_xlsx(rep):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"[:31]
    m = rep["meta"]
    ws.append([m["report"]])
    ws["A1"].font = Font(bold=True, size=14)
    for k in ("generated_at", "last_sync", "filters", "unresolved_warnings", "sensitive"):
        ws.append([k.replace("_", " ").title(), str(m[k])])
    for n in m["notes"]:
        ws.append(["Note", n])
    ws.append([])
    hdr = ws.max_row + 1
    ws.append([c["label"] for c in rep["columns"]])
    for c in ws[hdr]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2A44")
    for r in rep["rows"]:
        ws.append([_safe(r.get(c["key"])) if not isinstance(r.get(c["key"]), (dict, list)) else str(r.get(c["key"])) for c in rep["columns"]])
    for i, c in enumerate(rep["columns"], start=1):
        ws.column_dimensions[ws.cell(hdr, i).column_letter].width = max(14, min(48, len(c["label"]) + 6))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
