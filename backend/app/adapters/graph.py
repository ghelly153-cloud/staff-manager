"""Microsoft Graph adapter - production connection to a workbook in OneDrive for Business / SharePoint.

Uses only documented v1.0 endpoints:
  GET   /drives/{d}/items/{i}                                   (eTag/cTag/lastModified = change token)
  GET   /drives/{d}/items/{i}/workbook/worksheets               (list sheets + visibility)
  GET   .../worksheets/{name}/usedRange(valuesOnly=true)        (values + formulas)
  PATCH .../worksheets/{name}/range(address='A1')               (scoped cell writes)
  POST  .../workbook/createSession | closeSession               (persisted editing session)
  POST  .../workbook/worksheets/add                             (support sheets)
  POST  /drives/{d}/items/{i}/copy                              (backup)
  POST/PATCH/DELETE /subscriptions                              (change notifications on /drives/{d}/root)

STATUS: written against Microsoft's documentation and exercised against a simulated Graph in the test-suite.
It has NOT been run against a real Microsoft tenant by the author - see docs/TEST_RESULTS.md.
"""
import time
from datetime import date, datetime, timedelta, timezone
from urllib.parse import quote

import httpx

from .base import Adapter, WorkbookError
from ..workbook.schema import serial_from_date

GRAPH = "https://graph.microsoft.com/v1.0"
MAX_SUB_MINUTES = 42300            # documented maximum for driveItem subscriptions (<30 days)
DATE_FORMAT = "dd-mmm-yyyy"


class GraphAdapter(Adapter):
    name = "graph"
    is_live_production = True

    def __init__(self, drive_id, item_id, token_provider, client=None, sleep=time.sleep, max_retries=4):
        self.drive, self.item = drive_id, item_id
        self.tokens = token_provider
        self.client = client or httpx.Client(timeout=30)
        self.sleep, self.max_retries = sleep, max_retries
        self._file_name = None
        self._session_id = None

    # ---------- plumbing ----------
    @property
    def _item(self):
        return f"/drives/{self.drive}/items/{self.item}"

    def _ws(self, sheet):
        return f"{self._item}/workbook/worksheets/{quote(sheet, safe='')}"

    def _req(self, method, path, json=None, params=None, extra_headers=None, expect=(200, 201, 202, 204)):
        url = path if path.startswith("http") else GRAPH + path
        refreshed = False
        last = None
        for attempt in range(self.max_retries + 1):
            try:
                headers = {"Authorization": f"Bearer {self.tokens.get_token()}", "Accept": "application/json"}
                if self._session_id:
                    headers["workbook-session-id"] = self._session_id
                if extra_headers:
                    headers.update(extra_headers)
                r = self.client.request(method, url, json=json, params=params, headers=headers)
            except WorkbookError:
                raise
            except httpx.TransportError as ex:
                last = WorkbookError("network", f"Network problem talking to Microsoft ({type(ex).__name__})")
                self.sleep(min(2 ** attempt, 8))
                continue
            if r.status_code in expect:
                return r.json() if r.content and "json" in r.headers.get("content-type", "") else {"_status": r.status_code, "_headers": dict(r.headers)}
            code = self._err_code(r)
            if r.status_code == 401 and not refreshed:
                refreshed = True
                self.tokens.invalidate()
                continue
            if r.status_code in (429, 503, 504) or (r.status_code == 423):
                ra = self._retry_after(r, attempt)
                last = WorkbookError("throttled" if r.status_code == 429 else ("locked" if r.status_code == 423 else "network"),
                                     "Microsoft asked us to slow down" if r.status_code == 429 else
                                     ("The workbook is locked by another editor right now" if r.status_code == 423 else "Microsoft service temporarily unavailable"),
                                     retry_after=ra)
                if attempt < self.max_retries:
                    self.sleep(min(ra, 60))
                    continue
                raise last
            raise self._map_error(r.status_code, code)
        raise last or WorkbookError("network", "Could not reach Microsoft")

    @staticmethod
    def _err_code(r):
        try:
            return (r.json().get("error") or {}).get("code", "")
        except Exception:
            return ""

    @staticmethod
    def _retry_after(r, attempt):
        try:
            return max(1, int(r.headers.get("Retry-After", 2 ** (attempt + 1))))
        except ValueError:
            return 2 ** (attempt + 1)

    @staticmethod
    def _map_error(status, code):
        if status == 401:
            return WorkbookError("auth", "Microsoft sign-in has expired. An admin must reconnect the workbook.")
        if status == 403:
            return WorkbookError("permission", "The connected account does not have permission to this workbook.", detail=code)
        if status == 404:
            return WorkbookError("not_found", "The workbook (or sheet) was not found. It may have been moved, renamed or deleted.", detail=code)
        if status == 409:
            return WorkbookError("conflict", "Microsoft reported an edit conflict. Please retry.", detail=code)
        return WorkbookError("unknown", f"Microsoft returned an error ({status} {code}).", detail=code)

    # ---------- reading ----------
    def describe(self):
        return {"adapter": "graph", "label": "Microsoft 365 workbook (live connection)", "file": self._file_name or "(not read yet)"}

    def get_version(self):
        j = self._req("GET", self._item, params={"$select": "name,eTag,cTag,lastModifiedDateTime,size"})
        self._file_name = j.get("name", self._file_name)
        return f"{j.get('cTag','')}|{j.get('eTag','')}|{j.get('lastModifiedDateTime','')}"

    def _sheet_list(self):
        return self._req("GET", f"{self._item}/workbook/worksheets", params={"$select": "id,name,visibility"}).get("value", [])

    @staticmethod
    def _pad(rng):
        r0, c0 = rng.get("rowIndex", 0), rng.get("columnIndex", 0)
        def pad(m):
            rows = [[None if v == "" else v for v in row] for row in (m or [])]
            rows = [[None] * c0 + row for row in rows]
            return [[] for _ in range(r0)] + rows
        v = pad(rng.get("values"))
        f = pad(rng.get("formulas"))
        return v, f

    def _read_sheet_obj(self, ws):
        rng = self._req("GET", self._ws(ws["name"]) + "/usedRange(valuesOnly=true)",
                        params={"$select": "address,values,formulas,rowIndex,columnIndex"})
        v, f = self._pad(rng)
        return {"values": v, "formulas": f, "hidden": ws.get("visibility", "Visible") != "Visible"}

    def read_all(self):
        sheets = {}
        for ws in self._sheet_list():
            sheets[ws["name"]] = self._read_sheet_obj(ws)
        return {"sheets": sheets}

    def read_sheet(self, name):
        for ws in self._sheet_list():
            if ws["name"] == name:
                return self._read_sheet_obj(ws)
        return None

    # ---------- writing ----------
    class _Sess:
        def __init__(self, a): self.a = a
        def __enter__(self):
            try:
                j = self.a._req("POST", self.a._item + "/workbook/createSession", json={"persistChanges": True})
                self.a._session_id = j.get("id")
            except WorkbookError as ex:
                if ex.kind in ("auth", "permission", "not_found"):
                    raise
                self.a._session_id = None       # editing without a session is allowed by Graph
            return self
        def __exit__(self, *exc):
            sid = self.a._session_id
            if sid:
                try:
                    self.a._req("POST", self.a._item + "/workbook/closeSession", json={})
                except WorkbookError:
                    pass
                self.a._session_id = None

    def _patch(self, sheet, address, values=None, formulas=None, fmt=None):
        body = {}
        if formulas is not None:
            body["formulas"] = formulas
        if values is not None:
            body["values"] = values
        if fmt is not None:
            body["numberFormat"] = fmt
        self._req("PATCH", self._ws(sheet) + f"/range(address='{address}')", json=body)

    @staticmethod
    def _enc(v):
        """-> (is_formula, json_value, is_date)"""
        if v is None:
            return False, "", False
        if isinstance(v, datetime):
            v = v.date()
        if isinstance(v, date):
            return False, serial_from_date(v), True
        if isinstance(v, str) and v.startswith("="):
            return True, v, False
        return False, v, False

    def write_cells(self, sheet, cells):
        self._check_fail()
        with self._Sess(self):
            for a1, v in cells.items():
                isf, jv, isd = self._enc(v)
                if isf:
                    self._patch(sheet, a1, formulas=[[jv]])
                else:
                    self._patch(sheet, a1, values=[[jv]], fmt=[[DATE_FORMAT]] if isd else None)

    def write_block(self, sheet, top_left, rows):
        self._check_fail()
        if not rows:
            return
        from openpyxl.utils import coordinate_to_tuple, get_column_letter
        r0, c0 = coordinate_to_tuple(top_left)
        w = max(len(r) for r in rows)
        rows = [[self._enc(v)[1] for v in (list(r) + [None] * (w - len(r)))] for r in rows]
        addr = f"{top_left}:{get_column_letter(c0 + w - 1)}{r0 + len(rows) - 1}"
        with self._Sess(self):
            self._patch(sheet, addr, values=rows)

    def append_rows(self, sheet, rows):
        self._check_fail()
        cur = self.read_sheet(sheet)
        if cur is None:
            raise WorkbookError("not_found", f"Sheet “{sheet}” not found")
        vals = cur["values"]
        last = 0
        for i, r in enumerate(vals):
            if any(v not in (None, "") for v in r):
                last = i + 1
        first = last + 1
        self.write_block(sheet, f"A{first}", rows)
        return first

    def ensure_sheet(self, name, headers):
        if any(ws["name"] == name for ws in self._sheet_list()):
            return False
        self._check_fail()
        with self._Sess(self):
            self._req("POST", self._item + "/workbook/worksheets/add", json={"name": name})
            from openpyxl.utils import get_column_letter
            # keep ISO dates / keys as TEXT so Excel does not reinterpret them
            self._patch(name, "A2:B5000", fmt="@")
            self._patch(name, f"A1:{get_column_letter(len(headers))}1", values=[list(headers)])
        return True

    def backup(self, label):
        stem = (self._file_name or "workbook").rsplit(".", 1)[0]
        name = f"{stem}__BACKUP__{label}__{datetime.now():%Y%m%d-%H%M%S}.xlsx"
        r = self._req("POST", self._item + "/copy", json={"name": name}, expect=(200, 202))
        return name

    # ---------- change notifications (webhooks) ----------
    def subscribe(self, notification_url, lifecycle_url, client_state, minutes=MAX_SUB_MINUTES - 60):
        exp = (datetime.now(timezone.utc) + timedelta(minutes=minutes)).strftime("%Y-%m-%dT%H:%M:%S.0000000Z")
        body = {"changeType": "updated", "notificationUrl": notification_url, "lifecycleNotificationUrl": lifecycle_url,
                "resource": f"/drives/{self.drive}/root", "expirationDateTime": exp, "clientState": client_state}
        return self._req("POST", "/subscriptions", json=body)

    def renew_subscription(self, sub_id, minutes=MAX_SUB_MINUTES - 60):
        exp = (datetime.now(timezone.utc) + timedelta(minutes=minutes)).strftime("%Y-%m-%dT%H:%M:%S.0000000Z")
        return self._req("PATCH", f"/subscriptions/{sub_id}", json={"expirationDateTime": exp})

    def delete_subscription(self, sub_id):
        return self._req("DELETE", f"/subscriptions/{sub_id}", expect=(200, 204))
