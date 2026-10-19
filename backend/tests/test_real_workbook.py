"""Acceptance checks against a COPY of the real workbook. Run: REAL_WORKBOOK=/path/to/file.xlsx pytest tests/test_real_workbook.py"""
import os, shutil, hashlib, pytest
from datetime import date, timedelta
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from tests.conftest import H
from app.config import Settings
from app.main import create_app

SRC = os.getenv("REAL_WORKBOOK")
pytestmark = pytest.mark.skipif(not SRC, reason="set REAL_WORKBOOK to run")
EXPECTED = {"Babulal": 92, "Bhagwat": 58, "Naval": 28, "Surendra Chhipma": 31, "Bharatbhai Ferawala": 14, "Tejsingh Sisodiya": 8,
            "Shaileshbhai Prajapati": 3, "Hireshbhai Lakhvara": 2, "Dharmendrabhai staff": 5, "Babulal Fafda - Stadium": 10}

@pytest.fixture
def real(tmp_path):
    p = str(tmp_path / "copy.xlsx"); shutil.copy(SRC, p)
    s = Settings(source="local", local_workbook_path=p, dev_auth=True, data_dir=str(tmp_path / "d"))
    app = create_app(s, run_background=False); app.state.sync.refresh(force=True)
    c = TestClient(app); c.post("/auth/dev-login", json={"role": "admin"}, headers=H); return c, p, app

def test_ten_groups_and_exact_counts(real):
    c, p, app = real; d = c.get("/api/dashboard").json()
    assert {g["name"]: g["summary"]["total"] for g in d["groups"]} == EXPECTED and d["summary"]["total"] == 251 and d["skipped_sheets"] == []

def test_known_pay_numbers_and_no_invented_values(real):
    c, p, app = real; s = c.get("/api/dashboard").json()["summary"]
    # Financial policy: no payable totals authorised, confirmed count is 0, totals are None
    assert s["pay"]["rate_missing_count"] == 241 and s["pay"]["confirmed_count"] == 0 and s["pay"]["confirmed_total"] is None
    assert s["pay"]["employees_with_source_rate"] == 10  # 10 employees have rate cells in Excel, but rates remain unapproved
    babulal = next(i for i in c.get("/api/employees", params={"group": "Babulal", "page_size": 5}).json()["items"])
    assert babulal["legacy_days"] == 30 and babulal["amount"] is None and babulal["pay_status"] == "rates_unauthorised"
    assert s["not_verified"] == 251 and s["active"] == 0 and s["departed"] == 0              # statuses never assumed

def test_flags_expected_real_issues(real):
    c, p, app = real; codes = c.get("/api/quality").json()["counts"]
    for k in ("TITLE_MISMATCH", "SERIAL_UNUSUAL", "GROUP_AMOUNT_UNUSUAL", "RATE_REQUIRED", "COUNT_MISMATCH"): assert k in codes, k
    titles = [i["message"] for i in c.get("/api/quality", params={"show": "all"}).json()["items"] if i["code"] == "TITLE_MISMATCH"]
    assert any("Hireshbhai Lakhvara" in t and "Shaileshbhai Prajapati" in t for t in titles)
    assert any(i["code"] == "COUNT_MISMATCH" and "98" in i["message"] and "92" in i["message"] for i in c.get("/api/quality", params={"show": "all"}).json()["items"])

def test_group_balance_matches_workbook_formula(real):
    c, p, app = real; g = c.get("/api/groups/Babulal").json()["amounts"]
    # Raw spreadsheet cached formula value is preserved for transparency, but no app payable debt is asserted
    assert g["balance"]["workbook_cached"] == 2998 and g["balance"]["app"] is None and g["advance"]["value"] == 1

def test_setup_preserves_every_existing_cell_and_formula(real):
    c, p, app = real; before = load_workbook(p)
    snap = {(ws.title, x.coordinate): x.value for ws in before for row in ws.iter_rows(max_col=16) for x in row if x.value is not None}
    r = c.post("/api/setup/run", json={"confirmed": True}, headers=H); assert r.status_code == 200 and r.json()["ids_assigned"] == 251
    after = load_workbook(p)
    for (t, a), v in snap.items(): assert after[t][a].value == v, (t, a)
    assert len(set(e["id"] for e in c.get("/api/employees", params={"status": "all", "page_size": 200}).json()["items"])) == 200
    d = (date.today() - timedelta(days=1)).isoformat(); e = c.get("/api/employees", params={"group": "Naval"}).json()["items"][0]
    assert c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": d, "status": "Present"}]}, headers=H).status_code == 200
    assert c.get("/api/dashboard", params={"date": d}).json()["summary"]["attendance"]["present"] == 1
