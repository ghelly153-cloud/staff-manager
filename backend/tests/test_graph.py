"""Graph adapter exercised against a SIMULATED Graph. Proves our request/response handling; does NOT prove Microsoft behaviour."""
import shutil, httpx, pytest
from datetime import date, timedelta
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from tests.fixtures import build
from tests.fake_graph import FakeGraph, FakeTokens
from tests.conftest import H
from app.adapters.graph import GraphAdapter
from app.adapters.base import WorkbookError
from app.config import Settings
from app.main import create_app

@pytest.fixture
def g(tmp_path):
    p = str(tmp_path / "wb.xlsx"); build(p)
    fg, tok = FakeGraph(p), FakeTokens()
    ad = GraphAdapter("D", "I", tok, client=httpx.Client(transport=fg.transport), sleep=lambda s: None)
    s = Settings(source="local", local_workbook_path=p, dev_auth=True, data_dir=str(tmp_path / "d"), webhook_client_state="secret-state")
    app = create_app(s, adapter=ad, run_background=False); app.state.sync.refresh(force=True)
    c = TestClient(app); c.post("/auth/dev-login", json={"role": "admin"}, headers=H)
    class E: pass
    e = E(); e.fg, e.tok, e.ad, e.c, e.app, e.path, e.sync = fg, tok, ad, c, app, p, app.state.sync
    return e

def test_read_via_graph_discovers_groups(g):
    assert g.app.state.svc.snap.groups == ["Group One", "Group Two"] and len(g.app.state.svc.snap.employees) == 6
    assert g.sync.status()["state"] == "connected"

def test_full_write_flow_through_graph_requests(g):
    assert g.c.post("/api/setup/run", json={"confirmed": True}, headers=H).status_code == 200
    e = g.c.get("/api/employees", params={"status": "all"}).json()["items"][0]
    d = (date.today() - timedelta(days=1)).isoformat()
    assert g.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": d, "status": "Present"}]}, headers=H).status_code == 200
    assert load_workbook(g.path)["_App_Attendance"]["E2"].value == "Present"
    assert g.c.patch(f"/api/employees/{e['id']}", json={"changes": {"location": "Hall B"}, "expected": {"location": "Hall A"}}, headers=H).status_code == 200
    assert load_workbook(g.path)["Group One"]["D3"].value == "Hall B"
    assert g.fg.sessions_open == 0, "editing sessions must always be closed"
    assert any(c[0] == "POST" and c[1].endswith("/copy") for c in g.fg.calls), "backup copy requested before setup"
    only_scoped = [c for c in g.fg.calls if c[0] == "PATCH"]; assert all("range(address=" in c[1] for c in only_scoped)   # scoped writes only

def test_throttle_retries_then_succeeds(g):
    g.fg.inject_next(429, {"Retry-After": "1"}, n=2)
    assert g.sync.refresh(force=True) and g.sync.status()["state"] == "connected"

def test_persistent_throttle_reports_failure_and_backs_off(g):
    g.fg.inject_next(429, {"Retry-After": "7"}, n=20)
    assert g.sync.refresh(force=True) is False
    st = g.sync.status(); assert st["state"] == "sync_failed" and "slow down" in st["last_error"]["message"] and g.sync.delay() >= 7

def test_401_triggers_token_refresh_once(g):
    g.fg.inject_next(401); g.sync.refresh(force=True); assert g.tok.invalidated == 1 and g.sync.status()["state"] == "connected"

def test_expired_auth_shows_auth_required_and_keeps_last_data(g):
    g.tok.fail = True; g.sync.refresh(force=True)
    assert g.sync.status()["state"] == "auth_required" and len(g.app.state.svc.snap.employees) == 6     # last good data retained

def test_403_404_mapped(g):
    for code, kind in ((403, "permission"), (404, "not_found")):
        g.fg.inject_next(code)
        with pytest.raises(WorkbookError) as ex: g.ad.get_version()
        assert ex.value.kind == kind

def test_write_failure_over_graph_not_reported_as_success(g):
    g.c.post("/api/setup/run", json={"confirmed": True}, headers=H)
    e = g.c.get("/api/employees", params={"status": "all"}).json()["items"][0]
    g.fg.inject_next(500, n=1)                        # first request of the write (createSession) fails -> proceeds w/o session
    g.fg.inject_next(403, n=30)                       # then permission denied on the write itself
    r = g.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": (date.today() - timedelta(days=1)).isoformat(), "status": "Present"}]}, headers=H)
    assert r.status_code == 502 and "Not saved" in r.json()["message"]

def test_change_detection_only_rereads_when_token_changes(g):
    n0 = len([c for c in g.fg.calls if "usedRange" in c[1]]); assert g.sync.refresh() is False
    assert len([c for c in g.fg.calls if "usedRange" in c[1]]) == n0          # cheap metadata check only
    wb = load_workbook(g.path); wb["Group One"]["B8"] = "Typed In Excel"; wb["Group One"]["D8"] = "Hall A"; wb.save(g.path)
    assert g.sync.refresh() is True and any(e["name"] == "Typed In Excel" for e in g.app.state.svc.snap.employees)

def test_webhook_validation_handshake_and_notification(g):
    r = g.c.post("/webhooks/graph?validationToken=abc123"); assert r.status_code == 200 and r.text == "abc123" and r.headers["content-type"].startswith("text/plain")
    bad = g.c.post("/webhooks/graph", json={"value": [{"clientState": "WRONG"}]}); assert bad.status_code == 202 and g.sync.webhook["last_notification"] is None
    ok = g.c.post("/webhooks/graph", json={"value": [{"clientState": "secret-state"}]}); assert ok.status_code == 202 and g.sync.webhook["last_notification"]

def test_subscription_create_and_renew_use_documented_shape(g):
    j = g.ad.subscribe("https://x/webhooks/graph", "https://x/webhooks/graph/lifecycle", "secret-state"); assert j["id"] == "sub1"
    assert g.ad.renew_subscription("sub1")["id"] == "sub1"
