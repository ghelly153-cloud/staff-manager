"""Synthetic workbook that reproduces the REAL workbook's structure and quirks (fake names only)."""
from openpyxl import Workbook
HDR = ['No.', 'Staff Name', 'Designation', 'Loctation', 'Staff', 'Sept Present', 'Octo Present', 'Total Present', 'Per Day Pay', 'Net Pay', '2026 Per Day', 'Sweet Kg.', 'Aadhar', 'Mobile', 'From Date', 'To Date']

def _sheet(wb, name, title, people, blanks, total=True, rate_head='2026 Per Day'):
    ws = wb.create_sheet(name)
    ws['B1'], ws['C1'] = title, 'Extra Staff Test'
    for i, h in enumerate(HDR, 1): ws.cell(2, i, h)
    ws['K2'] = rate_head
    r = 3
    for p in people + [None] * blanks:
        ws.cell(r, 1, '=ROW()-ROW($A$2)'); ws.cell(r, 8, f'=SUM(F{r}:G{r})'); ws.cell(r, 10, f'=SUM(H{r}*I{r})')
        if p:
            for k, v in p.items(): ws.cell(r, {'name':2,'des':3,'loc':4,'staff':5,'sept':6,'octo':7,'rate':9,'mobile':14,'aadhaar':13}[k], v)
        r += 1
    if total:
        t = r + 2
        ws.cell(t, 1, f'=COUNT(A3:A{r-1})'); ws.cell(t, 2, 'Total'); ws.cell(t, 6, f'=SUM(F3:F{r-1})'); ws.cell(t, 8, f'=SUM(H3:H{r-1})'); ws.cell(t, 10, f'=SUM(J3:J{r-1})')
        ws.cell(t+1, 9, 'Advance'); ws.cell(t+1, 10, 1); ws.cell(t+2, 9, 'Sweets'); ws.cell(t+2, 10, 1)
        ws.cell(t+3, 9, 'Balance'); ws.cell(t+3, 10, f'=SUM(J{t}-J{t+1}-J{t+2})')
    return ws

def build(path):
    wb = Workbook(); wb.remove(wb.active)
    A = [dict(name='Test Person 01', des='Karigar', loc='Hall A', staff=1, sept=5, octo=25, rate=100),
         dict(name='Test Person 02', des='Helper', loc='Hall A', staff=1),
         dict(name='Test Person 03 ', des='Helper', loc='Hall B', staff=1, sept=0, octo=7.5),
         dict(name='Test Person 04', staff=1, octo=10, rate=100)]
    _sheet(wb, 'Group One', 'Group One', A, 6)
    B = [dict(name='Test Person 05', des='Cook', loc='Hall B', staff=1, octo=16),
         dict(name='test person 02', des='Cook', loc='Hall B', staff=1, octo=3)]   # duplicate-looking name
    _sheet(wb, 'Group Two', 'Totally Different Title', B, 2, total=False)
    ws = wb.create_sheet('Notes'); ws['A1'] = 'not an employee sheet'
    wb.save(path)
