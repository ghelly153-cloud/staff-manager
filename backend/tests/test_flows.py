"""D. App-to-Excel writes, conflicts, failures, status/departure, attendance, setup, new-row detection (Excel->app)."""
from datetime import date, timedelta
from openpyxl import load_workbook
from tests.conftest import H
from app.adapters.base import WorkbookError

TODAY = date.today()
def setup(env):
    r = env.c.post("/api/setup/run", json={"confirmed": True}, headers=H); assert r.status_code == 200, r.text; return r.json()
def emps(env, **q): return env.c.get("/api/employees", params={"page_size": 100, **q}).json()
def first(env, name="Test Person 01"): return next(i for i in emps(env, status="all")["items"] if i["name"] == name)

def test_read_only_until_setup(env):
    ref = emps(env)["items"][0]["ref"]
    r = env.c.post("/api/attendance", json={"items": [{"ref": ref, "date": TODAY.isoformat(), "status": "Present"}]}, headers=H)
    assert r.status_code == 409 and r.json()["code"] in ("id_setup_required",)

def test_setup_backup_ids_and_source_unchanged(env):
    before = load_workbook(env.path); snapshot = {(ws.title, c.coordinate): c.value for ws in before for row in ws.iter_rows(max_col=16) for c in row if c.value is not None}
    res = setup(env); assert res["ids_assigned"] == 6 and res["backup"]
    after = load_workbook(env.path)
    for (t, a), v in snapshot.items(): assert after[t][a].value == v, f"{t}!{a} changed"          # nothing existing modified
    assert {"_App_Attendance", "_App_Audit", "_App_Config"} <= set(after.sheetnames)
    ids = [e["id"] for e in emps(env, status="all")["items"]]; assert len(set(ids)) == 6 and all(i.startswith("EMP-") for i in ids)
    setup2 = env.c.post("/api/setup/run", json={"confirmed": True}, headers=H).json(); assert setup2["ids_assigned"] == 0    # idempotent, no duplicates
    assert [e["id"] for e in emps(env, status="all")["items"]] == ids

def test_attendance_write_verify_summary_and_duplicate_prevention(env):
    setup(env); e = first(env); d = (TODAY - timedelta(days=1)).isoformat()
    r = env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": d, "status": "Present"}]}, headers=H); assert r.status_code == 200, r.text
    ws = load_workbook(env.path)["_App_Attendance"]; rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == 1 and rows[0][2] == e["id"] and rows[0][4] == "Present"                        # really saved in the workbook
    env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": d, "status": "Absent", "expected": "Present"}]}, headers=H)
    ws = load_workbook(env.path)["_App_Attendance"]; assert ws.max_row == 2 and ws["E2"].value == "Absent"   # one row per employee+date
    s = env.c.get("/api/dashboard", params={"date": d}).json()["summary"]["attendance"]; assert s["available"] and s["absent"] == 1 and s["present"] == 0
    assert any(a["action"] == "attendance" for a in env.c.get("/api/audit").json()["items"])

def test_attendance_conflict_and_validation(env):
    setup(env); e = first(env); d = (TODAY - timedelta(days=1)).isoformat()
    env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": d, "status": "Present"}]}, headers=H)
    r = env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": d, "status": "Absent", "expected": "Not Recorded"}]}, headers=H)
    assert r.status_code == 409 and r.json()["code"] == "conflict"
    fut = (TODAY + timedelta(days=2)).isoformat()
    assert env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": fut, "status": "Present"}]}, headers=H).status_code == 422
    assert env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": d, "status": "Leave"}]}, headers=H).status_code == 422   # not enabled
    two = [{"ref": e["id"], "date": d, "status": "Present"}] * 2
    assert env.c.post("/api/attendance", json={"items": two}, headers=H).status_code == 422                                          # needs confirmation
    assert env.c.post("/api/attendance", json={"items": two, "confirmed": True}, headers=H).status_code == 422                       # duplicate in request

def test_failed_write_never_reports_success(env):
    setup(env); e = first(env); d = (TODAY - timedelta(days=1)).isoformat()
    env.app.state.adapter.fail_writes = True
    r = env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": d, "status": "Present"}]}, headers=H)
    assert r.status_code == 502 and r.json()["code"] == "write_failed" and "Not saved" in r.json()["message"]
    env.app.state.adapter.fail_writes = False
    assert env.c.get("/api/attendance", params={"date": d}).json()["summary"]["attendance"]["available"] is False
    assert env.sync.status()["state"] in ("sync_failed",)                                           # state shows the failure

def test_status_departure_flow_and_history_kept(env):
    setup(env); e = first(env); d1 = (TODAY - timedelta(days=3)).isoformat()
    env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": d1, "status": "Present"}]}, headers=H)
    assert env.c.post("/api/status", json={"items": [{"ref": e["id"], "status": "Departed"}]}, headers=H).status_code == 422          # date never guessed
    dep = (TODAY - timedelta(days=2)).isoformat()
    assert env.c.post("/api/status", json={"items": [{"ref": e["id"], "status": "Departed", "dep_date": dep, "reason": "Went home"}]}, headers=H).status_code == 200
    assert e["id"] not in [i["id"] for i in emps(env)["items"]]                                    # gone from default list
    dl = env.c.get("/api/departed").json(); assert dl["total"] == 1 and dl["items"][0]["last_att"] == d1 and dl["items"][0]["recent"]
    assert env.c.get("/api/dashboard").json()["summary"]["departed"] == 1
    assert e["id"] not in [i["id"] for i in env.c.get("/api/attendance", params={"date": d1}).json()["items"]]
    late = (TODAY - timedelta(days=1)).isoformat()
    assert env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": late, "status": "Present"}]}, headers=H).status_code == 422
    env.login("editor"); assert env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": late, "status": "Present"}], "correction": True}, headers=H).status_code == 422
    env.login("admin"); assert env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": late, "status": "Present"}], "correction": True}, headers=H).status_code == 200
    assert any(i["code"] == "ATTENDANCE_AFTER_DEPARTURE" for i in env.c.get("/api/quality").json()["items"])                        # flagged
    prof = env.c.get(f"/api/employees/{e['id']}").json(); assert prof["attendance_history"] and prof["status"] == "Departed"          # history retained
    ws = load_workbook(env.path)["Group One"]; assert ws["R3"].value == "Departed" and ws["B3"].value == "Test Person 01"           # row not deleted

def test_edit_persists_and_conflict_when_excel_changes(env):
    setup(env); e = first(env)
    r = env.c.patch(f"/api/employees/{e['id']}", json={"changes": {"location": "Hall B"}, "expected": {"location": "Hall A"}}, headers=H); assert r.status_code == 200, r.text
    assert load_workbook(env.path)["Group One"]["D3"].value == "Hall B"
    wb = load_workbook(env.path); wb["Group One"]["D3"] = "Hall A"; wb.save(env.path)            # someone edits in Excel meanwhile
    r = env.c.patch(f"/api/employees/{e['id']}", json={"changes": {"location": "Hall C"}, "expected": {"location": "Hall B"}}, headers=H)
    assert r.status_code == 409 and r.json()["code"] == "conflict"
    assert load_workbook(env.path)["Group One"]["D3"].value == "Hall A"                           # not silently overwritten
    env.login("editor"); assert env.c.patch(f"/api/employees/{e['id']}", json={"changes": {"rate": 100}, "expected": {}}, headers=H).status_code == 403
    env.login("admin"); assert env.c.patch(f"/api/employees/{e['id']}", json={"changes": {"mobile": "12345"}, "expected": {}}, headers=H).status_code == 422

def test_excel_to_app_new_row_edit_and_new_sheet(env):
    setup(env); n0 = emps(env, status="all")["total"]
    wb = load_workbook(env.path); ws = wb["Group One"]; ws["B8"], ws["C8"], ws["D8"] = "Test Person NEW", "Helper", "Hall A"; wb.save(env.path)     # new employee typed in Excel
    assert env.c.post("/api/sync/refresh", headers=H).status_code == 200
    items = emps(env, status="all"); assert items["total"] == n0 + 1
    new = next(i for i in items["items"] if i["name"] == "Test Person NEW"); assert new["id"] and new["id"].startswith("EMP-")      # auto ID, no duplicates
    n = env.c.get("/api/groups/Group One").json()["summary"]["total"]; assert n == 5
    wb = load_workbook(env.path); wb["Group One"]["B3"] = "Test Person 01 Corrected"; wb.save(env.path); env.sync.refresh(force=True)
    assert any(i["name"] == "Test Person 01 Corrected" for i in emps(env, status="all")["items"])
    from tests.fixtures import _sheet
    wb = load_workbook(env.path); _sheet(wb, "Group Three", "Group Three", [dict(name="Test Person 99", des="X", loc="Y", staff=1)], 1, total=False); wb.save(env.path)
    env.sync.refresh(force=True); assert "Group Three" in env.c.get("/api/bootstrap").json()["groups"]               # new group button appears
    wb = load_workbook(env.path); del wb["Group Three"]; wb.save(env.path); env.sync.refresh(force=True)
    assert "Group Three" not in env.c.get("/api/bootstrap").json()["groups"]

def test_row_move_keeps_identity(env):
    setup(env); e = first(env); wb = load_workbook(env.path); ws = wb["Group One"]
    for col in range(1, 22): a, b = ws.cell(3, col), ws.cell(5, col); ws.cell(3, col).value, ws.cell(5, col).value = (b.value if col not in (1, 8, 10) else a.value), (a.value if col not in (1, 8, 10) else b.value)
    wb.save(env.path); env.c.post("/api/sync/refresh", headers=H)
    p = env.c.get(f"/api/employees/{e['id']}").json(); assert p["name"] == "Test Person 01" and p["source"]["row"] == 5

def test_add_employee_with_duplicate_guard(env):
    setup(env)
    r = env.c.post("/api/employees", json={"data": {"group": "Group One", "name": "Brand New Person", "designation": "Helper", "location": "Hall A"}}, headers=H); assert r.status_code == 200, r.text
    ws = load_workbook(env.path)["Group One"]; row = r.json()["row"]; assert ws.cell(row, 2).value == "Brand New Person" and str(ws.cell(row, 8).value).startswith("=SUM(F")
    r2 = env.c.post("/api/employees", json={"data": {"group": "Group One", "name": "brand  new person"}}, headers=H); assert r2.status_code == 409 and r2.json()["code"] == "duplicate_candidate"
    env.login("editor"); r3 = env.c.post("/api/employees", json={"data": {"group": "Group One", "name": "Another One", "location": "Invented Hall"}}, headers=H); assert r3.status_code == 422

def test_config_rules_gate_datewise_pay(env):
    setup(env); e = first(env); d = (TODAY - timedelta(days=1)).isoformat()
    env.c.post("/api/attendance", json={"items": [{"ref": e["id"], "date": d, "status": "Present"}]}, headers=H)
    p = env.c.get("/api/payroll").json(); row = next(r for r in p["rows"] if r["id"] == e["id"]); assert row["pay_status"] == "rules_required" and row["amount"] is None
    sm = env.c.get("/api/payroll").json()["summary"]["pay"]
    assert sm["confirmed_total"] is None and sm["datewise_pending_rule_count"] == 1  # Financial policy: no payable totals authorised
    assert env.c.put("/api/config/legacy_cutoff_date", json={"value": (TODAY - timedelta(days=5)).isoformat()}, headers=H).status_code == 200
    assert env.c.put("/api/config/daily_rate_source", json={"value": "per_day_pay"}, headers=H).status_code == 200
    row = next(r for r in env.c.get("/api/payroll").json()["rows"] if r["id"] == e["id"])
    assert row["amount"] is None and row["pay_status"] == "rates_unauthorised" and row["datewise_days"] == 1  # 1 day recorded, unapproved rate
    env.login("editor"); assert env.c.put("/api/config/daily_rate_source", json={"value": None}, headers=H).status_code == 403

def test_export_has_metadata_and_csv_injection_guard(env):
    setup(env); e = first(env)
    assert env.c.patch(f"/api/employees/{e['id']}", json={"changes": {"name": "=HYPERLINK(1)"}, "expected": {}}, headers=H).status_code == 422   # never write formulas into the master workbook
    wb = load_workbook(env.path); wb["Group One"]["B4"] = "-2+3 odd"; wb.save(env.path); env.sync.refresh(force=True)
    body = env.c.get("/api/reports/all_employees?format=csv").text
    assert "# Generated:" in body and "# Last synced:" in body and "'-2+3 odd" in body
