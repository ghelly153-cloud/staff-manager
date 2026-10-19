"""A SIMULATED Microsoft Graph (httpx.MockTransport) backed by an .xlsx file. It is NOT Microsoft - it only lets us
exercise GraphAdapter's real request/response code, retries, sessions and error mapping offline."""
import re, os, json, httpx
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

class FakeGraph:
    def __init__(self, path, drive="D", item="I"):
        self.path, self.drive, self.item = path, drive, item
        self.calls, self.inject, self.sessions_open, self.subs = [], [], 0, {}
        self.transport = httpx.MockTransport(self.handle)

    def inject_next(self, status, headers=None, n=1):
        self.inject += [(status, headers or {})] * n

    def _j(self, obj, status=200, headers=None):
        return httpx.Response(status, json=obj, headers=headers or {})

    def handle(self, req: httpx.Request):
        path, m = req.url.path, req.method
        self.calls.append((m, path))
        if self.inject:
            st, hd = self.inject.pop(0)
            return httpx.Response(st, json={"error": {"code": "injected"}}, headers=hd)
        assert req.headers.get("authorization", "").startswith("Bearer "), "missing bearer token"
        base = f"/v1.0/drives/{self.drive}/items/{self.item}"
        if path == "/v1.0/subscriptions" and m == "POST":
            b = json.loads(req.content); sid = f"sub{len(self.subs)+1}"; self.subs[sid] = b
            assert b["changeType"] == "updated" and b["resource"] == f"/drives/{self.drive}/root"
            return self._j({"id": sid, "expirationDateTime": b["expirationDateTime"]}, 201)
        if path.startswith("/v1.0/subscriptions/") and m == "PATCH":
            return self._j({"id": path.rsplit("/", 1)[1], "expirationDateTime": json.loads(req.content)["expirationDateTime"]})
        if not path.startswith(base): return self._j({"error": {"code": "itemNotFound"}}, 404)
        rest = path[len(base):]
        if rest == "" and m == "GET":
            st = os.stat(self.path); return self._j({"name": "Test.xlsx", "eTag": f'"{st.st_mtime_ns}"', "cTag": f'"c{st.st_mtime_ns}"', "lastModifiedDateTime": "x", "size": st.st_size})
        if rest == "/copy": return httpx.Response(202, headers={"location": "https://x/monitor"})
        if rest == "/workbook/createSession": self.sessions_open += 1; return self._j({"id": "sess1", "persistChanges": True})
        if rest == "/workbook/closeSession": self.sessions_open -= 1; return httpx.Response(204)
        if rest == "/workbook/worksheets" and m == "GET":
            wb = load_workbook(self.path); return self._j({"value": [{"id": f"{{{i}}}", "name": ws.title, "visibility": "Visible" if ws.sheet_state == "visible" else "Hidden"} for i, ws in enumerate(wb.worksheets)]})
        if rest == "/workbook/worksheets/add":
            wb = load_workbook(self.path); wb.create_sheet(json.loads(req.content)["name"]); wb.save(self.path); return self._j({}, 201)
        mm = re.match(r"/workbook/worksheets/([^/]+)/(usedRange\(valuesOnly=true\)|range\(address='([^']+)'\))$", rest)
        if mm:
            name = mm.group(1); wb = load_workbook(self.path)
            if name not in wb.sheetnames: return self._j({"error": {"code": "ItemNotFound"}}, 404)
            ws = wb[name]
            if mm.group(2).startswith("usedRange"):
                rows = [[("" if v is None else v) for v in r] for r in ws.iter_rows(values_only=True)]
                vals = [["" if (isinstance(v, str) and v.startswith("=")) else v for v in r] for r in rows]
                return self._j({"address": f"{name}!A1", "rowIndex": 0, "columnIndex": 0, "values": vals, "formulas": rows})
            body = json.loads(req.content); a = mm.group(3)
            assert req.headers.get("workbook-session-id") == "sess1", "write outside a session"
            c1, r1, c2, r2 = range_boundaries(a)
            data = body.get("formulas") or body.get("values")
            if data is not None:
                for i, row in enumerate(data):
                    for j, v in enumerate(row): ws.cell(r1 + i, c1 + j).value = None if v == "" else v
                wb.save(self.path)
            return self._j({"address": a})
        return self._j({"error": {"code": "notImplemented", "path": rest}}, 400)

class FakeTokens:
    def __init__(self): self.invalidated = 0; self.fail = False
    def get_token(self):
        from app.adapters.base import WorkbookError
        if self.fail: raise WorkbookError("auth", "expired")
        return "fake-token"
    def invalidate(self): self.invalidated += 1
